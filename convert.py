"""Excel -> data.json çevirici.

Kullanım:
    pip install openpyxl
    python convert.py malzemeler.xlsx [data.json]

Excel'in ilk satırı başlık olmalı:
Malzeme No | Malzeme Açıklaması | Eldeki Miktar | Lokasyon No
(U/S/R No gibi başka sütunlar varsa yok sayılır.)

Çıktıya güncelleme tarihi/saati de yazılır (Türkiye saati).
"""
import json
import sys
from datetime import datetime, timedelta, timezone
from openpyxl import load_workbook

COLS = ["Malzeme No", "Malzeme Açıklaması", "Eldeki Miktar", "Lokasyon No"]


def temiz(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def main(dosya, cikti="data.json"):
    ws = load_workbook(dosya, data_only=True, read_only=True).active
    satirlar = ws.iter_rows(values_only=True)
    baslik = [temiz(h) for h in next(satirlar)]
    eksik = [c for c in COLS if c not in baslik]
    if eksik:
        sys.exit("Excel'de bu sütunlar bulunamadı: " + ", ".join(eksik))
    idx = {c: baslik.index(c) for c in COLS}

    veri = {}
    for s in satirlar:
        kod = temiz(s[idx["Malzeme No"]])
        if not kod:
            continue
        veri[kod] = {
            "aciklama": temiz(s[idx["Malzeme Açıklaması"]]),
            "miktar": temiz(s[idx["Eldeki Miktar"]]),
            "lokasyon": temiz(s[idx["Lokasyon No"]]),
        }

    simdi = datetime.now(timezone(timedelta(hours=3)))  # Türkiye saati (UTC+3)
    paket = {"guncelleme": simdi.strftime("%d.%m.%Y %H:%M"), "veri": veri}
    with open(cikti, "w", encoding="utf-8") as f:
        json.dump(paket, f, ensure_ascii=False)
    print(f"{len(veri)} malzeme {cikti} dosyasına yazıldı ({paket['guncelleme']}).")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        sys.exit("Kullanım: python convert.py malzemeler.xlsx [data.json]")
    main(*sys.argv[1:])
